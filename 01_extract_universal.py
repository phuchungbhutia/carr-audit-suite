"""
01_extract_universal.py
=======================
Stage 1: Multi-Year Adaptive Ingestion & Normalization.
Extracts Appendix II (Gram Panchayat Units) and Appendix III (Zilla Panchayats)
using validated physical page boundaries for CARR 2023, 2024, 2025, and 2026.
"""

from __future__ import annotations

import datetime
import os
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import pdfplumber

INPUT_DIR = Path("input_pdfs")
OUTPUT_DIR = Path("output")
OUTPUT_MASTER = OUTPUT_DIR / "carr_multiyear_master.xlsx"

# Exact validated 1-indexed page ranges for all four report volumes
REPORT_CONFIGS = [
    {
        "filename": "CARR 2023.pdf",
        "fy": "2021-22",
        "app2_start": 56,
        "app2_end": 72,
        "app3_start": 73,
        "app3_end": 73,
    },
    {
        "filename": "CARR 2024.pdf",
        "fy": "2022-23",
        "app2_start": 94,
        "app2_end": 122,
        "app3_start": 123,
        "app3_end": 123,
    },
    {
        "filename": "CARR 2025.pdf",
        "fy": "2023-24",
        "app2_start": 73,
        "app2_end": 91,
        "app3_start": 92,
        "app3_end": 93,
    },
    {
        "filename": "CARR 2026.pdf",
        "fy": "2024-25",
        "app2_start": 83,
        "app2_end": 107,
        "app3_start": 108,
        "app3_end": 108,
    },
]

DISTRICT_NAMES = [
    "GANGTOK", "PAKYONG", "GYALSHING", "MANGAN",
    "NAMCHI", "SORENG", "EAST", "WEST", "NORTH", "SOUTH"
]


def log_msg(level: str, msg: str) -> None:
    ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    icons = {"INFO": "[INFO]", "WARN": "[WARN]", "ERROR": "[FAIL]", "SUCCESS": "[ OK ]"}
    print(f"[{ts}] {icons.get(level, '[INFO]')} {msg}")


def parse_indian_amount(val: Any) -> float:
    if val is None or pd.isna(val):
        return 0.0
    s = str(val).strip()
    if not s or s.lower() in ["-", "--", "---", "nil", "null", "none", "n/a", ".."]:
        return 0.0

    is_negative = False
    if s.startswith("(") and s.endswith(")"):
        is_negative = True
        s = s[1:-1].strip()
    elif s.startswith("-"):
        is_negative = True
        s = s[1:].strip()

    s_clean = re.sub(r"[Rs₹$\s,]", "", s)
    try:
        f_val = float(s_clean)
        return -f_val if is_negative else f_val
    except ValueError:
        return 0.0


def normalize_scheme(raw_scheme: str) -> str:
    u = str(raw_scheme).upper()
    if "15" in u and ("FC" in u or "FINANCE" in u):
        return "15th Finance Commission"
    if "14" in u and ("FC" in u or "FINANCE" in u):
        return "14th Finance Commission"
    if "SFC" in u or "STATE FINANCE" in u:
        return "State Finance Commission"
    if "OSR" in u or "REVENUE" in u or "TRADE" in u or "LICENSE" in u:
        return "Own Source Revenue"
    if "GIA" in u or "SALARY" in u or "HONORARIUM" in u:
        return "Grant-In-Aid / Salaries"
    if "SBM" in u or "SWACHH" in u:
        return "Swachh Bharat Mission"
    if "JJM" in u:
        return "Jal Jeevan Mission"
    if "EFF" in u or "EMPOWERED" in u:
        return "Empowered Functionaries Fund"
    return "Other / Miscellaneous"


def extract_gpu_records(
    pdf: pdfplumber.PDF, start_page: int, end_page: int, fy: str, filename: str
) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    current_bac = "General"
    current_gpu = "General"

    for p in range(start_page, end_page + 1):
        if p > len(pdf.pages):
            break
        page = pdf.pages[p - 1]
        tables = page.extract_tables()
        for tbl in tables:
            for row in tbl:
                row_cells = [str(c).strip() if c is not None else "" for c in row]
                if len(row_cells) < 4:
                    continue

                full_txt = " ".join(row_cells).upper()
                if "OPENING BALANCE" in full_txt or "GRAND TOTAL" in full_txt:
                    continue

                # Filter subtotal rows
                if "TOTAL" in full_txt and not any(
                    x in full_txt for x in ["RECEIPT", "PAYMENT", "INTEREST"]
                ):
                    continue

                trailing_nums: List[Tuple[int, float]] = []
                for i in range(len(row_cells) - 1, -1, -1):
                    val = row_cells[i]
                    if val != "" and (re.search(r"\d", val) or val in ["-", "NIL", "nil"]):
                        parsed = parse_indian_amount(val)
                        trailing_nums.insert(0, (i, parsed))
                    else:
                        break

                if len(trailing_nums) < 3:
                    continue

                # Distinguish between 6-col (4 nums) and 7/8-col (5 nums) layouts
                if len(trailing_nums) >= 5:
                    ob = trailing_nums[-5][1]
                    tot_rec = trailing_nums[-3][1]
                    pay = trailing_nums[-2][1]
                    cb = trailing_nums[-1][1]
                    text_cutoff = trailing_nums[-5][0]
                elif len(trailing_nums) == 4:
                    ob = trailing_nums[-4][1]
                    tot_rec = trailing_nums[-3][1]
                    pay = trailing_nums[-2][1]
                    cb = trailing_nums[-1][1]
                    text_cutoff = trailing_nums[-4][0]
                else:
                    tot_rec = trailing_nums[-3][1]
                    pay = trailing_nums[-2][1]
                    cb = trailing_nums[-1][1]
                    ob = 0.0
                    text_cutoff = trailing_nums[-3][0]

                text_tokens = [c for c in row_cells[:text_cutoff] if c]
                if not text_tokens:
                    continue

                if len(text_tokens) >= 3:
                    current_bac = text_tokens[0]
                    current_gpu = text_tokens[1]
                    raw_scheme = " ".join(text_tokens[2:])
                elif len(text_tokens) == 2:
                    if any(
                        s in text_tokens[1].upper()
                        for s in ["FC", "SFC", "OSR", "GIA", "JJM", "SCHEME", "TOTAL"]
                    ):
                        current_gpu = text_tokens[0]
                        raw_scheme = text_tokens[1]
                    else:
                        current_bac = text_tokens[0]
                        current_gpu = text_tokens[1]
                        raw_scheme = "General"
                else:
                    raw_scheme = text_tokens[0]

                if tot_rec > 0 or pay > 0 or cb > 0 or ob > 0:
                    records.append({
                        "Financial_Year": fy,
                        "Tier": "Gram Panchayat Unit",
                        "Block_Administrative_Centre": current_bac.strip(),
                        "Entity_Name": current_gpu.strip(),
                        "Raw_Scheme": raw_scheme.strip(),
                        "Normalized_Bucket": normalize_scheme(raw_scheme),
                        "Opening_Balance": ob,
                        "Total_Receipt": tot_rec,
                        "Payment_Expenditure": pay,
                        "Closing_Balance": cb,
                        "Source_File": filename,
                        "Page_Number": p,
                    })

    return records


def extract_zp_records(
    pdf: pdfplumber.PDF, start_page: int, end_page: int, fy: str, filename: str
) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    current_zp = "District Zilla Panchayat"

    for p in range(start_page, end_page + 1):
        if p > len(pdf.pages):
            break
        page = pdf.pages[p - 1]
        tables = page.extract_tables()
        for tbl in tables:
            for row in tbl:
                row_cells = [str(c).strip() if c is not None else "" for c in row]
                if len(row_cells) < 4:
                    continue

                full_txt = " ".join(row_cells).upper()
                if "OPENING BALANCE" in full_txt or "GRAND TOTAL" in full_txt:
                    continue
                if "TOTAL" in row_cells[0].upper() or (len(row_cells) > 1 and "TOTAL" in row_cells[1].upper()):
                    continue

                nums = [parse_indian_amount(c) for c in row_cells if re.search(r"\d", c)]
                if len(nums) < 3:
                    continue

                for dist in DISTRICT_NAMES:
                    if dist in full_txt:
                        current_zp = f"{dist.title()} Zilla Panchayat"
                        break

                cb = nums[-1]
                pay = nums[-2]
                tot_rec = nums[-3]
                ob = nums[-4] if len(nums) >= 4 else 0.0
                raw_scheme = row_cells[1] if len(row_cells) > 1 and not re.search(r"\d", row_cells[1]) else "General"

                records.append({
                    "Financial_Year": fy,
                    "Tier": "Zilla Panchayat",
                    "Block_Administrative_Centre": "District Level",
                    "Entity_Name": current_zp,
                    "Raw_Scheme": raw_scheme,
                    "Normalized_Bucket": normalize_scheme(raw_scheme),
                    "Opening_Balance": ob,
                    "Total_Receipt": tot_rec,
                    "Payment_Expenditure": pay,
                    "Closing_Balance": cb,
                    "Source_File": filename,
                    "Page_Number": p,
                })

    return records


def format_excel(filepath: Path) -> None:
    wb = openpyxl.load_workbook(filepath)
    hdr_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    hdr_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=9)
    thin_border = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )

    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                cell.border = thin_border
                if cell.row > 1:
                    cell.font = data_font
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##,##0.00"
                        cell.alignment = Alignment(horizontal="right")
                val_len = len(str(cell.value or ""))
                if val_len > max_len:
                    max_len = val_len
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)

    wb.save(filepath)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    all_gpu: List[Dict[str, Any]] = []
    all_zp: List[Dict[str, Any]] = []

    log_msg("INFO", "Initializing Stage 1: Adaptive PDF Extraction...")

    for cfg in REPORT_CONFIGS:
        pdf_path = INPUT_DIR / cfg["filename"]
        if not pdf_path.exists():
            log_msg("WARN", f"File '{cfg['filename']}' not found in {INPUT_DIR.resolve()}. Skipping.")
            continue

        log_msg("INFO", f"Processing {cfg['filename']} for FY {cfg['fy']}...")
        with pdfplumber.open(pdf_path) as pdf:
            # GPU Extraction (Appendix II)
            extracted_gpus = extract_gpu_records(
                pdf, cfg["app2_start"], cfg["app2_end"], cfg["fy"], cfg["filename"]
            )
            all_gpu.extend(extracted_gpus)
            log_msg("INFO", f"  Extracted {len(extracted_gpus)} GPU rows from pages {cfg['app2_start']}-{cfg['app2_end']}.")

            # ZP Extraction (Appendix III)
            extracted_zps = extract_zp_records(
                pdf, cfg["app3_start"], cfg["app3_end"], cfg["fy"], cfg["filename"]
            )
            all_zp.extend(extracted_zps)
            log_msg("INFO", f"  Extracted {len(extracted_zps)} ZP rows from pages {cfg['app3_start']}-{cfg['app3_end']}.")

    if not all_gpu and not all_zp:
        log_msg("ERROR", "No records extracted. Ensure source PDFs are placed in 'input_pdfs/'.")
        return

    df_gpu = pd.DataFrame(all_gpu)
    df_zp = pd.DataFrame(all_zp)

    with pd.ExcelWriter(OUTPUT_MASTER, engine="openpyxl") as writer:
        if not df_gpu.empty:
            df_gpu.to_excel(writer, sheet_name="GPU_Master_Ledger", index=False)
        if not df_zp.empty:
            df_zp.to_excel(writer, sheet_name="ZP_Master_Ledger", index=False)

    format_excel(OUTPUT_MASTER)
    log_msg("SUCCESS", f"Master extraction complete. Saved to: {OUTPUT_MASTER.resolve()}")


if __name__ == "__main__":
    main()